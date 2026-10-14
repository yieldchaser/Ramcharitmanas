# -*- coding: utf-8 -*-
"""Build a formatted .xlsx and .docx track list from the merge log + saved durations."""
import os, re, sys, json, datetime
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

LOG = r"C:\Users\Dell\AppData\Local\Temp\ytdl\merge_music.log"
DUR = r"C:\Users\Dell\AppData\Local\Temp\ytdl\bare_durations.json"
WS = r"C:\Users\Dell\Github\Ramcharitmanas"
XLSX = os.path.join(WS, "Music-Tracklist-1-399.xlsx")
DOCX = os.path.join(WS, "Music-Tracklist-1-399.docx")

GROUPS = ["Chaiti", "Kajri", "Jhumar", "Hori", "Phag", "Rajan Ji Maharaj",
          "Rambhadracharya", "Ramnandacharya", "Swami Raghvacharya", "Madhushala"]
GROUP_FILL = {  # soft pastel cell fills
    "Chaiti": "FCE4D6", "Kajri": "FFF2CC", "Jhumar": "E2EFDA", "Hori": "DDEBF7",
    "Phag": "EAD1DC", "Rajan Ji Maharaj": "DAEEF3", "Rambhadracharya": "E4DFEC",
    "Ramnandacharya": "D6EFE3", "Swami Raghvacharya": "FDEADA",
    "Madhushala": "D6DCE4",
}
NAVY, OLD_FILL, NEW_FILL, ZEBRA = "1F4E79", "EDF3FA", "EBF6EC", "F7F9FB"

# ---------- rebuild the rows (same source as the .txt list) ----------
copied, renamed = {}, {}
for line in open(LOG, encoding="utf-8", errors="replace"):
    line = line.rstrip("\n")
    m = re.match(r"^copied (\d+): (.+?)  \(\d+ KB\)$", line)
    if m:
        copied[int(m.group(1))] = m.group(2)
        continue
    m = re.match(r"^renamed (\d+)\.mp3 -> (.+)$", line)
    if m:
        renamed[int(m.group(1))] = m.group(2)

durations = json.load(open(DUR, encoding="utf-8"))
rows, missing = [], []
for n in range(1, 400):
    if n >= 278:
        name, old = copied[n], False
    elif n in renamed:
        name, old = renamed[n], False
    else:
        name, old = f"{n:02d}.mp3", True
    group = ""
    if n >= 278:
        tail = name[:-4].rsplit(" - ", 1)[-1]
        group = tail if tail in GROUPS else ""
    raw = f"{n:02d}.mp3"
    secs = durations.get(raw, durations.get(f"{n}.mp3")) if old and " - " not in name else None
    rows.append({"n": n, "name": name, "group": group, "secs": secs, "old": old})

if len(copied) != 122 or len(renamed) != 121 or len(rows) != 399:
    missing = [len(copied), len(renamed), len(rows)]
    print("PARSE PROBLEM:", missing); sys.exit(1)
untitled = [r for r in rows if r["old"] and " - " not in r["name"]]
if len(untitled) != 156:
    print("PARSE PROBLEM untitled:", len(untitled)); sys.exit(1)
print("rows ok: 399 (122 new, 121 renamed, 156 untitled)")

def hms(s):
    if s is None:
        return "—", None
    h, rem = divmod(int(s), 3600); m, sec = divmod(rem, 60)
    txt = f"{h}:{m:02d}:{sec:02d}" if h else f"{m}:{sec:02d}"
    return txt, (h * 3600 + m * 60 + sec) / 86400.0     # fraction of a day for Excel

# group number ranges
ranges = {}
for r in rows:
    if r["group"]:
        a = ranges.setdefault(r["group"], [r["n"], r["n"], 0])
        a[1], a[2] = r["n"], a[2] + 1

# ================= Excel =================
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.properties import PageSetupProperties

thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()
ws = wb.active
ws.title = "Track List"

ws.merge_cells("A1:E1")
c = ws["A1"]
c.value = "MUSIC FOLDER  —  COMPLETE TRACK LIST"
c.font = Font(bold=True, size=18, color="FFFFFF")
c.fill = PatternFill("solid", fgColor=NAVY)
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 36

ws.merge_cells("A2:E2")
c = ws["A2"]
c.value = ("Numbers 1–399   ·   1–277 old collection   ·   278–399 new downloads   ·   "
           "Length is only known where the file itself records it (the untitled ones)")
c.font = Font(italic=True, size=10, color="404040")
c.fill = PatternFill("solid", fgColor="DCE6F1")
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 20
ws.row_dimensions[3].height = 6

HEADERS = ["No.", "File name", "Group", "Length", "Type"]
WIDTHS = [6, 88, 24, 11, 9]
for i, (h, w) in enumerate(zip(HEADERS, WIDTHS), start=1):
    cell = ws.cell(row=4, column=i, value=h)
    cell.font = Font(bold=True, size=11, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2E75B6")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = BORDER
    ws.column_dimensions[chr(64 + i)].width = w
ws.row_dimensions[4].height = 24

for idx, r in enumerate(rows):
    row = 5 + idx
    zebra = ZEBRA if idx % 2 else "FFFFFF"
    txt, frac = hms(r["secs"])
    vals = [r["n"], r["name"], r["group"] or "—", txt, "Old" if r["old"] else "New"]
    for col, v in enumerate(vals, start=1):
        cell = ws.cell(row=row, column=col, value=v)
        cell.border = BORDER
        cell.fill = PatternFill("solid", fgColor=zebra)
        cell.alignment = Alignment(vertical="center",
                                   horizontal="center" if col in (1, 4, 5) else "left",
                                   wrap_text=(col == 2))
    ws.cell(row=row, column=1).font = Font(bold=True, size=10)
    ws.cell(row=row, column=5).fill = PatternFill("solid",
        fgColor=OLD_FILL if r["old"] else NEW_FILL)
    if r["group"]:
        ws.cell(row=row, column=3).fill = PatternFill("solid", fgColor=GROUP_FILL[r["group"]])
        ws.cell(row=row, column=3).font = Font(bold=True, size=10)
    if frac is not None:
        cc = ws.cell(row=row, column=4)
        cc.number_format = "mm:ss" if r["secs"] < 3600 else "[h]:mm:ss"
    if r["secs"] is None:
        ws.cell(row=row, column=4).font = Font(color="A6A6A6", italic=True)

last = 4 + len(rows)
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:E{last}"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
ws.print_title_rows = "4:4"
ws.oddFooter.center.text = "Page &P of &N"
ws.oddFooter.right.text = "Music folder track list"
ws.sheet_view.showGridLines = False

# ---- sheet 2: index / how to read ----
ix = wb.create_sheet("Index & notes")
ix.merge_cells("A1:D1")
c = ix["A1"]
c.value = "GROUP INDEX  —  which numbers belong to which folder"
c.font = Font(bold=True, size=15, color="FFFFFF")
c.fill = PatternFill("solid", fgColor=NAVY)
c.alignment = Alignment(horizontal="center", vertical="center")
ix.row_dimensions[1].height = 30
for i, (h, w) in enumerate(zip(["Group", "Numbers", "Count", "Colour"], [30, 22, 10, 12]), start=1):
    cell = ix.cell(row=3, column=i, value=h)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2E75B6")
    cell.alignment = Alignment(horizontal="center")
    cell.border = BORDER
    ix.column_dimensions[chr(64 + i)].width = w
for i, g in enumerate(GROUPS):
    lo, hi, cnt = ranges[g]
    rr = 4 + i
    for col, v in enumerate([g, f"{lo}–{hi}", cnt, ""], start=1):
        cell = ix.cell(row=rr, column=col, value=v)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="left" if col in (1, 2) else "center")
    ix.cell(row=rr, column=1).fill = PatternFill("solid", fgColor=GROUP_FILL[g])
    ix.cell(row=rr, column=1).font = Font(bold=True)
    ix.cell(row=rr, column=3).fill = PatternFill("solid", fgColor=GROUP_FILL[g])

start = 4 + len(GROUPS) + 2
notes = [
    ("OLD COLLECTION (1–277)", None),
    ("named from inside the file", "121 files"),
    ("untitled — no song info inside, only length", "156 files"),
    ("", ""),
    ("HOW TO READ THIS LIST", None),
    ("Say the number out loud to point at a song", "e.g. “number 302”"),
    ("Old files keep their original number (01–277)", "untitled ones show length to identify them"),
    ("New downloads run 278–399", "group at the end of the name = its folder"),
    ("Type column", "Old = was already on the drive, New = downloaded in this project"),
]
ix.cell(row=start - 1, column=1, value="NOTES").font = Font(bold=True, size=12, color=NAVY)
for i, (a, b) in enumerate(notes):
    rr = start + i
    ca = ix.cell(row=rr, column=1, value=a)
    if b is None and a:
        ca.font = Font(bold=True, color=NAVY)
    ix.cell(row=rr, column=2, value=b or "")
ix.column_dimensions["A"].width = 52
ix.column_dimensions["B"].width = 46
ix.sheet_view.showGridLines = False

wb.save(XLSX)
print("wrote", XLSX, os.path.getsize(XLSX), "bytes")

# ================= Word =================
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)

def set_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Inches(11), Inches(8.5)
for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, attr, Inches(0.6))
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10)
rpr = style.element.get_or_add_rPr()
rf = rpr.find(qn("w:rFonts"))
if rf is None:
    rf = OxmlElement("w:rFonts"); rpr.append(rf)
rf.set(qn("w:ascii"), "Calibri"); rf.set(qn("w:hAnsi"), "Calibri"); rf.set(qn("w:cs"), "Nirmala UI")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("MUSIC FOLDER — COMPLETE TRACK LIST")
run.bold = True; run.font.size = Pt(24); run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Numbers 1 to 399   ·   277 old songs + 122 new downloads")
run.italic = True; run.font.size = Pt(12); run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

p = doc.add_paragraph()
p.add_run("How it is organised.  ").bold = True
p.add_run("Every line is numbered and the number is the thing to say out loud. "
          "Numbers 1–277 are the old collection: 121 of them were renamed from the title "
          "stored inside the file, and 156 carry no song info at all — those show only "
          "their length, so they can be identified by ear and renamed later. "
          "Numbers 278–399 are the new downloads; the last word in each name is the "
          "folder/group it came from. The Length column is filled in only where the "
          "file itself records a duration.")

p = doc.add_paragraph()
p.add_run("Group index.  ").bold = True
p.add_run("Where each folder sits in the number range:")

gi = doc.add_table(rows=1, cols=4)
gi.style = "Table Grid"
gi.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = ["Group", "Numbers", "Count"]
for i, h in enumerate(hdr):
    cell = gi.rows[0].cells[i]
    cell.text = ""
    run = cell.paragraphs[0].add_run(h)
    run.bold = True; run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); run.font.size = Pt(10)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    shade(cell, "2E75B6")
for g in GROUPS:
    lo, hi, cnt = ranges[g]
    cells = gi.add_row().cells
    cells[0].text = g; cells[1].text = f"{lo}–{hi}"; cells[2].text = str(cnt)
    for i, cell in enumerate(cells):
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i else WD_ALIGN_PARAGRAPH.LEFT
        for r_ in cell.paragraphs[0].runs:
            r_.font.size = Pt(10)
        shade(cell, GROUP_FILL[g] if i < 3 else "FFFFFF")
    cells[0].paragraphs[0].runs[0].bold = True
set_widths(gi, [Inches(3.0), Inches(2.2), Inches(1.2)])

doc.add_page_break()
p = doc.add_paragraph()
run = p.add_run("THE LIST")
run.bold = True; run.font.size = Pt(16); run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

WIDTHS = [Inches(0.55), Inches(6.35), Inches(2.0), Inches(0.9)]
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["No.", "File name", "Group", "Length"]
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = ""
    run = cell.paragraphs[0].add_run(h)
    run.bold = True; run.font.size = Pt(10); run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    shade(cell, "2E75B6")
# repeat header row on every page
trPr = table.rows[0]._tr.get_or_add_trPr()
th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)

for idx, r in enumerate(rows):
    cells = table.add_row().cells
    txt, _ = hms(r["secs"])
    vals = [str(r["n"]), r["name"], r["group"] or "—", txt]
    zebra = "F7F9FB" if idx % 2 else "FFFFFF"
    for i, (cell, v) in enumerate(zip(cells, vals)):
        cell.text = ""
        run = cell.paragraphs[0].add_run(v)
        run.font.size = Pt(8.5)
        cell.paragraphs[0].alignment = (WD_ALIGN_PARAGRAPH.CENTER if i in (0, 3)
                                       else WD_ALIGN_PARAGRAPH.LEFT)
        if i == 0:
            run.bold = True
        shade(cell, zebra)
    if r["group"]:
        shade(cells[2], GROUP_FILL[r["group"]])
        cells[2].paragraphs[0].runs[0].bold = True
    if r["secs"] is None:
        cells[3].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xA6, 0xA6, 0xA6)
        cells[3].paragraphs[0].runs[0].italic = True
set_widths(table, WIDTHS)

# ---- page-number footer (so printed pages are numbered) ----
def add_field(par, instr, cached):
    r1 = par.add_run()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r1._r.append(fc)
    r2 = par.add_run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    r2._r.append(it)
    r3 = par.add_run()
    fs = OxmlElement("w:fldChar"); fs.set(qn("w:fldCharType"), "separate"); r3._r.append(fs)
    r4 = par.add_run(cached)
    r5 = par.add_run()
    fe = OxmlElement("w:fldChar"); fe.set(qn("w:fldCharType"), "end"); r5._r.append(fe)
    return [r1, r2, r3, r4, r5]

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
lead = fp.add_run("Page ")
lead.font.size = Pt(9); lead.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
for r_ in add_field(fp, " PAGE ", "1"):
    r_.font.size = Pt(9); r_.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
mid = fp.add_run(" of ")
mid.font.size = Pt(9); mid.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
for r_ in add_field(fp, " NUMPAGES ", "1"):
    r_.font.size = Pt(9); r_.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

core = doc.core_properties
core.title = "Music Folder Track List 1-399"
core.author = "Track list generator"
doc.save(DOCX)
print("wrote", DOCX, os.path.getsize(DOCX), "bytes")
print("DONE")
